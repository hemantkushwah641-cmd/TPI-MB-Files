"""Long-lived master Excel across sessions. Tracks every portal bill and what was downloaded."""
from __future__ import annotations

import os
import re
from datetime import datetime
from pathlib import Path

from urllib.parse import quote

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

HEADERS = [
    "Scheme ID",
    "MB No.",
    "LoA No.",
    "TPI Received Date",
    "TPI Officer",
    "Cluster",
    "District Name",
    "Contractor Name",
    "Scheme Name",
    "Grand Total",
    "Measurement Date",
    "Bill Type",
    "Unit Name",
    "DL List",
    "DL LOA",
    "DL Comments",
    "DL MB Files",
    "DL Docs",
    "Last session",
    "Last updated",
    "Folder link",
]

FLAG_MAP = {
    "loa": "DL LOA",
    "comments": "DL Comments",
    "mb": "DL MB Files",
    "docs": "DL Docs",
    "excel": "DL List",
}


def master_path() -> Path:
    p = (os.getenv("TPI_MASTER") or "").strip()
    if p:
        return Path(p)
    root = Path(os.getenv("DOWNLOAD_DIR") or "downloads")
    return root / "tpi_master.xlsx"


def _norm_mb(mb: str) -> str:
    return (mb or "").strip().upper().replace("-", "/").replace(" ", "")


def _norm(s: str) -> str:
    return " ".join(str(s or "").split()).strip().upper()


def track_key(row: dict) -> str:
    sid = _norm(row.get("scheme_id") or row.get("Scheme ID"))
    mb = _norm_mb(row.get("mb_no") or row.get("MB No."))
    loa = _norm(row.get("loa_number") or row.get("list_ref") or row.get("LoA No."))
    tpi = _norm(row.get("tpi_level_date") or row.get("TPI Received Date"))
    return f"{sid}|{mb}|{loa}|{tpi}"


def _onedrive_mounts() -> list[tuple[Path, str]]:
    """Windows: local OneDrive folder → SharePoint/OneDrive web root."""
    out: list[tuple[Path, str]] = []
    try:
        import winreg
    except ImportError:
        return out

    def _enum(root, path):
        try:
            key = winreg.OpenKey(root, path)
        except OSError:
            return
        i = 0
        while True:
            try:
                subname = winreg.EnumKey(key, i)
            except OSError:
                break
            i += 1
            try:
                sub = winreg.OpenKey(key, subname)
            except OSError:
                continue

            def gv(*names):
                for n in names:
                    try:
                        v = winreg.QueryValueEx(sub, n)[0]
                        if v:
                            return str(v)
                    except OSError:
                        continue
                return ""

            mount = gv("MountPoint", "UserFolder", "DisplayNamePath")
            url = gv("UrlNamespace", "Url", "WebPath", "SharePointUrl", "LibraryUrl")
            if mount and url and url.lower().startswith("http"):
                out.append((Path(mount), url.rstrip("/")))
            try:
                winreg.CloseKey(sub)
            except OSError:
                pass
        try:
            winreg.CloseKey(key)
        except OSError:
            pass

    _enum(winreg.HKEY_CURRENT_USER, r"Software\SyncEngines\Providers\OneDrive")
    for acct in ("Business1", "Business2", "Personal"):
        _enum(winreg.HKEY_CURRENT_USER, rf"Software\Microsoft\OneDrive\Accounts\{acct}")
    return out


def kautilya_data_url(folder: str) -> str:
    """https:// OneDrive/SharePoint URL only — never C:\\ or file://."""
    folder = str(folder or "").strip()
    if not folder:
        return ""
    if folder.lower().startswith("https://"):
        return folder
    if folder.lower().startswith("http://") or folder.lower().startswith("file:"):
        return ""
    try:
        local = Path(folder).resolve()
    except Exception:
        return ""

    web = (os.getenv("TPI_ONEDRIVE_WEB") or "").strip().rstrip("/")
    root = (os.getenv("DOWNLOAD_DIR") or "").strip()
    if web.lower().startswith("https://") and root:
        try:
            rel = local.relative_to(Path(root).resolve())
            return web + "/" + "/".join(quote(str(p)) for p in rel.parts)
        except Exception:
            pass

    for mount, url in _onedrive_mounts():
        try:
            rel = local.relative_to(mount.resolve())
        except Exception:
            continue
        return url.rstrip("/") + "/" + "/".join(quote(str(p)) for p in rel.parts)
    return ""


def snapshot_name(session: str = "") -> str:
    sess = (session or os.getenv("TPI_SESSION") or "Session").strip() or "Session"
    sess = re.sub(r'[<>:"/\\|?*]', "-", sess)
    stamp = datetime.now().strftime("%Y-%m-%d %H%M")
    return f"tpi_master_{sess}_{stamp}.xlsx"


def _row_from_excel(headers: list, raw: tuple) -> dict:
    d = {}
    for i, h in enumerate(headers):
        d[h] = raw[i] if i < len(raw) and raw[i] is not None else ""
    return d


def load_master(path: Path | None = None) -> list[dict]:
    path = Path(path or master_path())
    if not path.exists():
        return []
    try:
        wb = load_workbook(path, data_only=False)
        ws = wb.active
        rows = list(ws.iter_rows())
    except Exception:
        return []
    if not rows:
        wb.close()
        return []
    headers = [str(c.value or "").strip() for c in rows[0]]
    try:
        link_i = headers.index("Folder link")
    except ValueError:
        link_i = -1
    out = []
    for raw in rows[1:]:
        if not raw or all(c.value is None or str(c.value).strip() == "" for c in raw):
            continue
        vals = tuple(c.value for c in raw)
        d = _row_from_excel(headers, vals)
        if link_i >= 0 and link_i < len(raw):
            cell = raw[link_i]
            href = ""
            if cell.hyperlink and getattr(cell.hyperlink, "target", None):
                href = str(cell.hyperlink.target)
            elif str(cell.value or "").lower().startswith("https://"):
                href = str(cell.value)
            if href.lower().startswith("https://"):
                d["Folder link"] = href
            elif str(d.get("Folder link") or "").lower() in ("kautilya data", "open folder"):
                d["Folder link"] = ""
        out.append(d)
    wb.close()
    return out


def save_master(rows: list[dict], path: Path | None = None) -> Path:
    path = Path(path or master_path())
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Master"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    link_i = HEADERS.index("Folder link") + 1
    n_ok = n_miss = 0
    for r in rows:
        ws.append([r.get(h, "") for h in HEADERS])
        folder = str(r.get("Folder link") or r.get("save_folder") or "").strip()
        href = kautilya_data_url(folder)
        cell = ws.cell(ws.max_row, link_i)
        if href.startswith("https://"):
            cell.value = "Kautilya Data"
            cell.hyperlink = href
            cell.font = Font(color="0563C1", underline="single")
            r["Folder link"] = href
            n_ok += 1
        else:
            cell.value = ""
            n_miss += 1
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"
    for col in ws.columns:
        width = min(max(len(str(c.value or "")) for c in col) + 2, 50)
        ws.column_dimensions[col[0].column_letter].width = width
    wb.save(path)
    try:
        if path.name.lower() == "tpi_master.xlsx":
            wb.save(path.with_name(snapshot_name()))
    except Exception:
        pass
    if n_miss:
        print(
            f"[{datetime.now().strftime('%H:%M:%S')}] Folder link: {n_ok} OneDrive https, {n_miss} skipped (no OneDrive web URL).",
            flush=True,
        )
    else:
        print(
            f"[{datetime.now().strftime('%H:%M:%S')}] Folder link: {n_ok} OneDrive https links.",
            flush=True,
        )
    return path


def find_record(master: list[dict], row: dict) -> dict | None:
    sid = _norm(row.get("scheme_id") or row.get("Scheme ID"))
    mb = _norm_mb(row.get("mb_no") or row.get("MB No."))
    loa = _norm(row.get("loa_number") or row.get("list_ref") or row.get("LoA No."))
    tpi = _norm(row.get("tpi_level_date") or row.get("TPI Received Date"))
    if not sid or not mb:
        return None
    exact = []
    loose = []
    for rec in master:
        if _norm(rec.get("Scheme ID")) != sid or _norm_mb(rec.get("MB No.")) != mb:
            continue
        rloa = _norm(rec.get("LoA No."))
        rtpi = _norm(rec.get("TPI Received Date"))
        if tpi and rtpi and tpi == rtpi and (not loa or not rloa or loa == rloa):
            exact.append(rec)
        elif loa and rloa and loa == rloa and (not tpi or not rtpi or tpi == rtpi):
            exact.append(rec)
        else:
            loose.append(rec)
    if exact:
        return exact[0]
    if len(loose) == 1:
        return loose[0]
    return None


def remaining_steps(rec: dict | None, wanted: set[str]) -> set[str]:
    open_steps = {"loa", "comments", "mb", "docs"}
    need = wanted & open_steps
    if rec is None:
        return need
    miss = set()
    for step, col in FLAG_MAP.items():
        if step in need and str(rec.get(col) or "").strip().upper() != "Y":
            miss.add(step)
    return miss


def from_portal_row(row: dict, session: str = "") -> dict:
    loa = row.get("loa_number") or row.get("list_ref") or ""
    return {
        "Scheme ID": str(row.get("scheme_id") or ""),
        "MB No.": str(row.get("mb_no") or ""),
        "LoA No.": str(loa),
        "TPI Received Date": str(row.get("tpi_level_date") or ""),
        "TPI Officer": str(row.get("tpi_officer") or ""),
        "Cluster": str(row.get("cluster") or os.getenv("TPI_CLUSTER") or ""),
        "District Name": str(row.get("district") or row.get("district_folder") or ""),
        "Contractor Name": str(row.get("contractor") or ""),
        "Scheme Name": str(row.get("scheme") or ""),
        "Grand Total": str(row.get("grand_total") or "").replace(",", ""),
        "Measurement Date": str(row.get("date") or row.get("measurement_date") or ""),
        "Bill Type": str(row.get("bill_type") or ""),
        "Unit Name": str(row.get("unit") or ""),
        "DL List": "Y",
        "DL LOA": "",
        "DL Comments": "",
        "DL MB Files": "",
        "DL Docs": "",
        "Last session": session,
        "Last updated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "Folder link": str(row.get("save_folder") or ""),
    }


def apply_flags(rec: dict, row: dict, done_steps: set[str], session: str) -> dict:
    fresh = from_portal_row(row, session)
    for h in HEADERS:
        val = fresh.get(h, "")
        if val not in ("", None):
            rec[h] = val
        elif h not in rec:
            rec[h] = ""
    rec["DL List"] = "Y"
    if row.get("tpi_level_date"):
        rec["TPI Received Date"] = row.get("tpi_level_date") or rec.get("TPI Received Date")
    if "loa" in done_steps or row.get("grand_total") or row.get("tpi_level_date"):
        rec["DL LOA"] = "Y"
    if "comments" in done_steps and row.get("comments_file"):
        rec["DL Comments"] = "Y"
    if "mb" in done_steps and (row.get("signed_pdf") or row.get("abstract") or row.get("signed_excel")):
        rec["DL MB Files"] = "Y"
    if "docs" in done_steps and row.get("uploaded_docs"):
        rec["DL Docs"] = "Y"
    rec["Last session"] = session or rec.get("Last session") or ""
    rec["Last updated"] = datetime.now().strftime("%Y-%m-%d %H:%M")
    return rec


def upsert_rows(master: list[dict], rows: list[dict], done_steps: set[str], session: str) -> list[dict]:
    for row in rows:
        rec = find_record(master, row)
        if rec is None:
            rec = from_portal_row(row, session)
            rec = apply_flags(rec, row, done_steps, session)
            master.append(rec)
        else:
            apply_flags(rec, row, done_steps, session)
    return master


def sync_gsheet(xlsx_path: Path, sheet_url: str) -> str:
    """Optional Google Sheet. Needs gspread + google_service_account.json next to the app."""
    url = (sheet_url or os.getenv("TPI_GSHEET") or "").strip()
    if not url:
        return ""
    cred = Path(__file__).resolve().parent / "google_service_account.json"
    try:
        import gspread
        from google.oauth2.service_account import Credentials
    except Exception:
        return "Google Sheets skipped (install gspread + google-auth, add google_service_account.json)"
    if not cred.exists():
        return f"Google Sheets skipped — put service account JSON at {cred.name}"
    scopes = [
        "https://www.googleapis.com/auth/spreadsheets",
        "https://www.googleapis.com/auth/drive",
    ]
    creds = Credentials.from_service_account_file(str(cred), scopes=scopes)
    gc = gspread.authorize(creds)
    sh = gc.open_by_url(url) if "http" in url else gc.open_by_key(url)
    try:
        ws = sh.worksheet("Master")
    except Exception:
        ws = sh.add_worksheet("Master", rows=2000, cols=24)
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    data = [[c if c is not None else "" for c in row] for row in wb.active.iter_rows(values_only=True)]
    wb.close()
    ws.clear()
    if data:
        ws.update("A1", data)
    return f"Google Sheet updated: {sh.url}"
