"""
TPI Measurement Book — Upload.

Puts files from each bill's TO_UPLOAD folder onto the open bill
on https://upjnr-kautilya.in (file inputs on the detail page).

Folder layout (same as download):
  Kautilya E-MB / DD.MM.YYYY / Session N / District / {scheme}_{...} / TO_UPLOAD / *

Usage:
  python tpi_uploader.py --headed --session "Session 1" --day 07.09.2026
  python tpi_uploader.py --headed --scan-only
"""
from __future__ import annotations

import argparse
import os
import re
import sys
import tempfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

from tpi_mb_downloader import (
    ACTION_DELAY,
    BASE_URL,
    DOWNLOAD_DIR,
    USER,
    PASSWORD,
    back_to_tpi_list,
    bill_folder_name,
    dismiss_popups,
    go_to_tpi_list,
    log,
    login,
    open_row,
    save_debug,
    scrape_grand_total,
    today_stamp,
    win_name,
)

from playwright.sync_api import sync_playwright


def parse_bill_folder(name: str) -> dict:
    m = re.match(r"^(\d{6,})_", name)
    sid = m.group(1) if m else ""
    ra = re.search(r"RA[-_/]?(CIVIL|EM)[-_/]?(\d+)", name, re.I)
    mb = ""
    if ra:
        kind = ra.group(1).upper()
        num = ra.group(2)
        mb = f"RA/{kind}/{num}" if kind == "CIVIL" else f"RA/EM/{num}"
    return {"scheme_id": sid, "mb_no": mb, "folder_name": name}


UPLOAD_DIR_NAMES = ("For Upload", "TO_UPLOAD")
MAX_BYTES = 7 * 1024 * 1024


def _norm_bill_type(text: str) -> str:
    t = (text or "").strip().lower()
    if t.startswith("ret") or t in ("r", "return bill"):
        return "return"
    return "forward"


def fill_from_excel(bill_dir: Path) -> dict:
    """Amount Approved, TPI Letter Number, Remark from xlsx in bill or For Upload."""
    out = {"amount": "", "letter": "", "remark": "", "bill_type": "forward"}
    paths = []
    for name in UPLOAD_DIR_NAMES:
        d = bill_dir / name
        if d.exists():
            paths.extend(d.glob("*.xlsx"))
            paths.extend(d.glob("*.xls"))
    paths.extend(bill_dir.glob("*.xlsx"))
    paths = [p for p in paths if p.is_file() and not p.name.startswith("~$")
             and "signed" not in p.name.lower() and "tpi_output" not in p.name.lower()]
    if not paths:
        return out
    from openpyxl import load_workbook
    for path in paths:
        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(max_row=30, max_col=12, values_only=True))
            wb.close()
        except Exception:
            continue
        if not rows:
            continue
        headers = [str(c or "").strip().lower() for c in rows[0]]

        def col(*keys):
            for i, h in enumerate(headers):
                if any(k in h for k in keys):
                    return i
            return None

        ia = col("amount approved", "amount")
        il = col("letter")
        ir = col("remark", "comment")
        it = col("bill type", "action")
        # vertical key-value sheet
        if ia is None and il is None and it is None:
            for row in rows:
                k = str(row[0] or "").strip().lower()
                v = row[1] if len(row) > 1 else ""
                if "amount" in k:
                    out["amount"] = str(v or "").strip()
                elif "letter" in k:
                    out["letter"] = str(v or "").strip()
                elif "remark" in k or "comment" in k:
                    out["remark"] = str(v or "").strip()
                elif "bill type" in k or k in ("type", "action", "forward", "return"):
                    out["bill_type"] = str(v or "").strip()
            if out["amount"] or out["letter"] or out["remark"] or out["bill_type"]:
                out["bill_type"] = _norm_bill_type(out["bill_type"])
                if out["bill_type"] == "return":
                    out["amount"] = "0.00"
                log(f"  fill excel: {path.name} type={out['bill_type']}")
                return out
            continue
        data = rows[1] if len(rows) > 1 else None
        if not data:
            continue
        if ia is not None:
            out["amount"] = str(data[ia] or "").strip()
        if il is not None:
            out["letter"] = str(data[il] or "").strip()
        if ir is not None:
            out["remark"] = str(data[ir] or "").strip()
        if it is not None:
            out["bill_type"] = str(data[it] or "").strip()
        out["bill_type"] = _norm_bill_type(out["bill_type"])
        if out["bill_type"] == "return":
            out["amount"] = "0.00"
        log(f"  fill excel: {path.name} type={out['bill_type']}")
        return out
    return out


def list_upload_files(bill_dir: Path) -> list[Path]:
    files: list[Path] = []
    for name in UPLOAD_DIR_NAMES:
        d = bill_dir / name
        if not d.exists():
            continue
        for p in sorted(d.iterdir()):
            if not p.is_file() or p.name.startswith("~$") or p.name.startswith("."):
                continue
            if p.suffix.lower() in {".xlsx", ".xls"}:
                continue
            files.append(p)
    return files


def scan_session(root: Path, day: str, session: str) -> list[dict]:
    sess = root / day / session
    if not sess.exists():
        log(f"Session folder not found: {sess}")
        return []
    bills = []
    for dist in sorted(p for p in sess.iterdir() if p.is_dir()):
        if dist.name.lower().startswith("session"):
            continue
        for bill in sorted(p for p in dist.iterdir() if p.is_dir()):
            meta = parse_bill_folder(bill.name)
            files = list_upload_files(bill)
            fill = fill_from_excel(bill)
            bills.append({
                **meta,
                "district": dist.name,
                "path": str(bill),
                "files": [str(f) for f in files],
                "n_files": len(files),
                "fill": fill,
            })
    return bills


def list_days(root: Path) -> list[str]:
    if not root.exists():
        return []
    days = [p.name for p in root.iterdir() if p.is_dir() and re.match(r"\d{2}\.\d{2}\.\d{4}$", p.name)]
    return sorted(days, reverse=True)


def list_sessions(root: Path, day: str) -> list[str]:
    d = root / day
    if not d.exists():
        return []
    sess = [p.name for p in d.iterdir() if p.is_dir() and re.search(r"^session\s+\d+", p.name, re.I)]
    return sorted(sess, key=lambda s: int(re.search(r"\d+", s).group(0)))


def upload_files_on_detail(page, files: list[str]) -> list[str]:
    """Choose File → Description → UPLOAD, one file at a time (max 7 MB)."""
    saved: list[str] = []
    try:
        page.get_by_text("Upload file if any", exact=False).first.scroll_into_view_if_needed(timeout=8000)
    except Exception:
        page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
    page.wait_for_timeout(500)

    file_input = page.locator("input[type='file']")
    if file_input.count() == 0:
        save_debug(page, "no_file_input")
        log("  No Choose File control on this page")
        return saved

    for path in files:
        p = Path(path)
        if not p.exists():
            continue
        if p.stat().st_size > MAX_BYTES:
            log(f"  ERROR: file over 7 MB, skipped: {p.name} ({p.stat().st_size/1024/1024:.1f} MB)")
            continue
        log(f"  uploading {p.name}")
        try:
            file_input.first.set_input_files(str(p))
            page.wait_for_timeout(400)
            desc = page.get_by_placeholder(re.compile("file description", re.I))
            if desc.count() == 0:
                desc = page.locator("input[placeholder*='Description' i], textarea[placeholder*='Description' i]")
            if desc.count():
                desc.first.fill(p.stem[:80])
            btn = page.get_by_role("button", name=re.compile(r"^\s*upload\s*$", re.I))
            if btn.count() == 0:
                btn = page.get_by_text(re.compile(r"^\s*upload\s*$", re.I))
            if btn.count():
                btn.first.click(timeout=8000)
            else:
                page.evaluate(
                    """() => {
                        const b = [...document.querySelectorAll('button,a,.btn')].find(e =>
                            /^\\s*upload\\s*$/i.test((e.innerText||'').replace(/\\s+/g,' ').trim())
                        );
                        if (b) b.click();
                    }"""
                )
            page.wait_for_timeout(800)
            click_modal_ok(page, wait_ms=4000)
            dismiss_popups(page, wait_ms=400)
            saved.append(str(p))
            log(f"  uploaded {p.name}")
        except Exception as exc:
            log(f"  upload fail {p.name}: {exc}")
            save_debug(page, f"upload_fail_{p.stem[:30]}")
    return saved


def _type_into(page, loc, value: str, tab: bool = True) -> bool:
    """Angular fields ignore paste — type like a keyboard."""
    value = str(value)
    try:
        if not loc.is_visible():
            log("    skip type - field not visible")
            return False
    except Exception:
        pass
    try:
        loc.scroll_into_view_if_needed(timeout=3000)
    except Exception:
        pass
    loc.click(timeout=5000)
    page.wait_for_timeout(200)
    loc.press("Control+A")
    page.wait_for_timeout(80)
    loc.press("Backspace")
    page.wait_for_timeout(80)
    page.keyboard.type(value, delay=70)
    page.wait_for_timeout(150)
    if tab:
        loc.press("Tab")
        page.wait_for_timeout(250)
    try:
        got = loc.input_value()
    except Exception:
        got = ""
    if str(got).strip() != value.strip():
        loc.click()
        loc.press("Control+A")
        loc.press("Backspace")
        page.keyboard.type(value, delay=90)
        if tab:
            loc.press("Tab")
            page.wait_for_timeout(200)
        try:
            got = loc.input_value()
        except Exception:
            got = ""
    log(f"    typed '{value}' now='{got}'")
    return str(got).strip() == value.strip() or (got.replace(",", "") == value.replace(",", ""))


def fill_tpi_fields(page, fill: dict) -> None:
    amount = (fill or {}).get("amount") or ""
    letter = (fill or {}).get("letter") or ""
    remark = (fill or {}).get("remark") or ""
    if not (amount or letter or remark):
        log("  no Amount / Letter / Remark in excel")
        return

    try:
        page.get_by_text("Amount Approved by TPI", exact=False).first.scroll_into_view_if_needed(timeout=8000)
    except Exception:
        pass
    page.wait_for_timeout(400)

    if amount:
        loc = page.locator(
            "xpath=//*[contains(normalize-space(.),'Amount Approved by TPI')][1]/following::input[1]"
        )
        if loc.count() == 0:
            loc = page.get_by_label(re.compile("Amount Approved by TPI", re.I))
        ok = False
        if loc.count():
            ok = _type_into(page, loc.first, amount)
        log(f"  Amount Approved by TPI = {amount} ({ok})")

    if letter:
        loc = page.get_by_placeholder(re.compile("TPI Letter Number", re.I))
        if loc.count() == 0:
            loc = page.locator(
                "xpath=//*[contains(normalize-space(.),'TPI Letter Number')][1]/following::input[1]"
            )
        ok = False
        if loc.count():
            ok = _type_into(page, loc.first, letter)
        log(f"  TPI Letter Number = {letter} ({ok})")

    if remark:
        ta = page.locator("textarea")
        ok = False
        try:
            if ta.count():
                box = ta.last
                box.scroll_into_view_if_needed(timeout=5000)
                box.click()
                box.fill("")
                page.keyboard.type(str(remark), delay=25)
                ok = True
        except Exception as exc:
            log(f"  remark type: {exc}")
        log(f"  Remark filled ({ok})")
    page.wait_for_timeout(500)


def click_modal_ok(page, wait_ms: int = 8000) -> str:
    """OK on Success / OTP sent / Are you sure want to return|forward."""
    page.wait_for_timeout(400)
    for _ in range(max(1, wait_ms // 250)):
        info = page.evaluate(
            """() => {
                const body = (document.body && document.body.innerText) || '';
                const vis = [...document.querySelectorAll('button, .btn, a')].filter(b => {
                    const r = b.getBoundingClientRect();
                    const t = (b.innerText || '').trim();
                    return r.width && r.height && /^OK$/i.test(t);
                });
                const hit = vis.find(b => {
                    const box = b.closest('.modal, .swal2-popup, [class*="modal"], [class*="alert"], [role="dialog"]') || b.parentElement;
                    const t = (box && box.innerText) || body;
                    return /success|data save successfully|please enter|tpi number|otp has been sent|are you sure|want to return|want to forward|press ok/i.test(t);
                }) || vis[0];
                if (!hit) return {ok: false};
                const msg = (hit.closest('.modal, .swal2-popup, [class*="modal"], [role="dialog"]') || hit.parentElement).innerText || '';
                hit.click();
                return {ok: true, msg: msg.replace(/\\s+/g, ' ').trim().slice(0, 160)};
            }"""
        )
        if info and info.get("ok"):
            log(f"  popup OK: {info.get('msg')}")
            page.wait_for_timeout(500)
            return str(info.get("msg") or "ok")
        page.wait_for_timeout(250)
    return ""


def click_save_as_draft(page) -> bool:
    btn = page.get_by_role("button", name=re.compile("save as draft", re.I))
    if btn.count() == 0:
        btn = page.get_by_text(re.compile("save as draft", re.I))
    try:
        if btn.count():
            btn.first.click(timeout=8000)
        else:
            page.evaluate(
                """() => {
                    const b = [...document.querySelectorAll('button,a,.btn')].find(e =>
                        /save as draft/i.test(e.innerText || '')
                    );
                    if (b) b.click();
                }"""
            )
        page.wait_for_timeout(800)
        msg = click_modal_ok(page, wait_ms=10000)
        if re.search(r"Please Enter TPI Number|Enter TPI (Letter )?Number", msg or "", re.I):
            log("  alert: Please Enter TPI Number — letter field empty")
            return False
        if re.search(r"save successfully|Success", msg or "", re.I):
            log("  SAVE AS DRAFT — Data save successfully")
            try:
                page.wait_for_load_state("domcontentloaded", timeout=8000)
            except Exception:
                pass
            page.wait_for_timeout(800)
            try:
                dismiss_popups(page)
            except Exception:
                pass
            return True
        # fallback: maybe modal already gone
        try:
            body = page.inner_text("body")[:4000]
        except Exception:
            body = ""
        if re.search(r"Please Enter TPI Number", body, re.I):
            click_modal_ok(page, wait_ms=2000)
            return False
        log("  clicked SAVE AS DRAFT")
        return True
    except Exception as exc:
        log(f"  SAVE AS DRAFT: {exc}")
        save_debug(page, "save_draft_fail")
        return False


def click_attach_signature(page) -> bool:
    """Forward bills only. DSC dongle/PIN is completed by the user if the OS dialog appears."""
    btn = page.get_by_role("button", name=re.compile("attach digital signature", re.I))
    if btn.count() == 0:
        btn = page.get_by_text(re.compile("attach digital signature", re.I))
    try:
        if btn.count():
            btn.first.click(timeout=8000)
        else:
            clicked = page.evaluate(
                """() => {
                    const b = [...document.querySelectorAll('button,a,.btn')].find(e =>
                        /attach digital signature/i.test(e.innerText || '')
                    );
                    if (b) { b.click(); return true; }
                    return false;
                }"""
            )
            if not clicked:
                log("  ATTACH DIGITAL SIGNATURE button not found")
                return False
        log("  clicked ATTACH DIGITAL SIGNATURE")
        # Alert: Are you sure you want to Attach DSC?  →  Yes, please!
        page.wait_for_timeout(600)
        yes = page.get_by_role("button", name=re.compile(r"yes,?\s*please", re.I))
        if yes.count() == 0:
            yes = page.get_by_text(re.compile(r"yes,?\s*please", re.I))
        if yes.count():
            yes.first.click(timeout=8000)
            log("  clicked Yes, please! on DSC confirm")
        else:
            page.evaluate(
                """() => {
                    const b = [...document.querySelectorAll('button,a,.btn')].find(e =>
                        /yes,?\s*please/i.test(e.innerText || '')
                    );
                    if (b) b.click();
                }"""
            )
            log("  DSC confirm Yes, please! (js)")
        page.wait_for_timeout(3000)
        click_modal_ok(page, wait_ms=4000)
        return True
    except Exception as exc:
        log(f"  DSC: {exc}")
        save_debug(page, "dsc_fail")
        return False


def settle_detail(page) -> None:
    """Wait until Save Draft reload finishes and the bill page is usable."""
    log("  waiting for bill page after Save Draft...")
    for _ in range(6):
        try:
            page.wait_for_load_state("domcontentloaded", timeout=4000)
            break
        except Exception as exc:
            if "destroyed" in str(exc).lower() or "navigation" in str(exc).lower():
                page.wait_for_timeout(600)
                continue
            break
    try:
        page.wait_for_url(re.compile(r"third-party-inspection/record"), timeout=12000)
    except Exception:
        pass
    try:
        dismiss_popups(page)
    except Exception:
        pass
    page.wait_for_timeout(600)


def _scroll_action_into_view(page) -> None:
    try:
        page.locator("xpath=//*[normalize-space()='Action:' or normalize-space()='Action']").last.scroll_into_view_if_needed()
    except Exception:
        try:
            page.get_by_text("Select One", exact=True).last.scroll_into_view_if_needed()
        except Exception:
            pass
    page.wait_for_timeout(250)


def _click_select_one(page) -> bool:
    _scroll_action_into_view(page)
    locators = (
        page.locator("xpath=//*[normalize-space()='Action:' or normalize-space()='Action']/following::*[contains(normalize-space(.),'Select One')][1]"),
        page.get_by_text("Select One", exact=True).last,
        page.locator("ng-select, .ng-select, [role='combobox']").last,
    )
    for attempt in range(5):
        for loc in locators:
            try:
                if not loc.count():
                    continue
                loc.first.scroll_into_view_if_needed()
                loc.first.click(timeout=3500)
                log("  clicked Select One (Action box)")
                return True
            except Exception as exc:
                s = str(exc).lower()
                if "destroyed" in s or "navigation" in s:
                    try:
                        page.wait_for_load_state("domcontentloaded", timeout=8000)
                    except Exception:
                        pass
                    page.wait_for_timeout(600)
                    break
                continue
        page.wait_for_timeout(400)
    log("  Select One not clickable")
    return False


def _native_action_select(page, want: str) -> str:
    try:
        n = page.locator("select").count()
    except Exception:
        return ""
    for i in range(n):
        try:
            sel = page.locator("select").nth(i)
            opts = [o.strip() for o in sel.locator("option").all_text_contents()]
            hit = next(
                (
                    o
                    for o in opts
                    if (want == "return" and re.search(r"return\s+to", o, re.I))
                    or (want == "forward" and re.search(r"forward\s+to", o, re.I))
                ),
                None,
            )
            if not hit:
                continue
            sel.scroll_into_view_if_needed()
            sel.select_option(label=hit)
            try:
                sel.dispatch_event("change")
                sel.dispatch_event("input")
            except Exception:
                pass
            log(f"  native Action select -> {hit}")
            return hit
        except Exception:
            continue
    return ""


def _click_overlay_option(page, want: str) -> bool:
    """Click Return/Forward inside the dropdown panel (even if it opens on the left)."""
    pat = r"Return To AE" if want == "return" else r"Forward To"
    page.wait_for_timeout(500)
    try:
        dump = page.evaluate(
            """() => {
                const panels = [...document.querySelectorAll(
                    '.cdk-overlay-pane, .ng-dropdown-panel, .ng-dropdown-panel-items, .dropdown-menu, .p-dropdown-items, [class*="overlay-pane"], [class*="dropdown-panel"]'
                )];
                return panels.slice(0, 6).map(p => {
                    const r = p.getBoundingClientRect();
                    return {cls: (p.className || '').toString().slice(0, 80),
                            t: (p.innerText || '').replace(/\\s+/g, ' ').trim().slice(0, 180),
                            x: Math.round(r.x), y: Math.round(r.y), w: Math.round(r.width), h: Math.round(r.height)};
                });
            }"""
        )
        log(f"  overlay panels: {dump}")
    except Exception as exc:
        log(f"  overlay dump: {exc}")
    panel = page.locator(
        ".cdk-overlay-pane, .ng-dropdown-panel, .ng-dropdown-panel-items, .dropdown-menu.show, .p-dropdown-items, [class*='overlay-pane']"
    )
    try:
        if panel.count():
            hit = panel.last.get_by_text(re.compile(pat, re.I))
            if hit.count():
                hit.last.click(timeout=4000, force=True)
                log(f"  overlay panel option force-clicked ({pat})")
                return True
    except Exception as exc:
        log(f"  overlay panel click: {exc}")
    try:
        info = page.evaluate(
            """(want) => {
                const norm = t => (t || '').replace(/\\s+/g, ' ').trim();
                const re = want === 'return' ? /return\\s+to\\s+ae/i : /forward\\s+to/i;
                const panels = [...document.querySelectorAll(
                    '.cdk-overlay-pane, .ng-dropdown-panel, .ng-dropdown-panel-items, .dropdown-menu, .p-dropdown-items, [class*="overlay-pane"], [class*="dropdown-panel"]'
                )];
                const roots = panels.length ? panels : [document.body];
                const hits = [];
                for (const root of roots) {
                    for (const e of root.querySelectorAll('div, span, li, a, p, option')) {
                        if (e.tagName === 'BUTTON' || e.closest('button')) continue;
                        const t = norm(e.innerText);
                        if (!re.test(t) || t.length > 90 || t.length < 8) continue;
                        const r = e.getBoundingClientRect();
                        if (r.width < 40 || r.height < 12 || r.height > 56) continue;
                        hits.push({t, x: r.x + r.width / 2, y: r.y + r.height / 2, h: r.height, w: r.width});
                    }
                }
                hits.sort((a, b) => a.t.length - b.t.length);
                return hits[0] || null;
            }""",
            want,
        )
        log(f"  overlay hit: {info}")
        if info and info.get("x") is not None:
            page.mouse.click(float(info["x"]), float(info["y"]))
            log(f"  mouse overlay -> {info.get('t')}")
            return True
    except Exception as exc:
        log(f"  overlay js: {exc}")
    return False


def _action_current_text(page) -> str:
    try:
        return page.evaluate(
            """() => {
                const norm = t => (t || '').replace(/\\s+/g, ' ').trim();
                const labs = [...document.querySelectorAll('label,div,span,p,strong')].filter(e => {
                    const t = norm(e.innerText);
                    const r = e.getBoundingClientRect();
                    return (t === 'Action:' || t === 'Action') && r.height > 0 && r.height < 40 && r.width < 220;
                });
                labs.sort((a, b) => b.getBoundingClientRect().y - a.getBoundingClientRect().y);
                const lab = labs[0];
                if (!lab) return '';
                const el = lab.nextElementSibling || lab.parentElement;
                return norm((el && el.innerText) || '');
            }"""
        ) or ""
    except Exception:
        return ""


def _otp_field_visible(page) -> bool:
    """True only if OTP input/label is actually on screen (not hidden Angular field)."""
    try:
        loc = page.get_by_placeholder(re.compile(r"otp", re.I))
        for i in range(loc.count()):
            el = loc.nth(i)
            if el.is_visible():
                box = el.bounding_box()
                if box and box["width"] > 40 and box["height"] > 10:
                    return True
        lab = page.get_by_text(re.compile(r"Please Enter OTP", re.I))
        for i in range(lab.count()):
            el = lab.nth(i)
            if el.is_visible():
                box = el.bounding_box()
                if box and box["width"] > 40:
                    return True
    except Exception:
        pass
    return False


def _pick_from_open_list(page, want: str) -> bool:
    """Click Return / Forward in the open Action list. No keyboard typing."""
    pat = r"Return To AE" if want == "return" else r"Forward To"
    page.wait_for_timeout(400)
    try:
        nodes = page.locator("div, li, span, a, p").filter(has_text=re.compile(pat, re.I))
        n = nodes.count()
        log(f"  list rows={n} for /{pat}/")
        for i in range(n):
            el = nodes.nth(i)
            try:
                if not el.is_visible():
                    continue
                box = el.bounding_box()
                if not box or box["height"] < 16 or box["height"] > 50 or box["width"] < 80:
                    continue
                cls = (el.get_attribute("class") or "") + " "
                if re.search(r"\bbtn\b|btn-danger", cls, re.I):
                    continue
                el.click(timeout=3000, force=True)
                log(f"  clicked list option i={i} h={int(box['height'])} w={int(box['width'])}")
                page.wait_for_timeout(400)
                return True
            except Exception as exc:
                log(f"  list click i={i}: {str(exc)[:120]}")
                continue
    except Exception as exc:
        log(f"  list scan: {exc}")
    return _click_overlay_option(page, want)


def select_action(page, btype: str) -> bool:
    """Click Select One under Action, then Return or Forward. Retries if page reloads."""
    want = "return" if btype == "return" else "forward"
    log(f"  Action: pick {want}")
    last = ""
    for attempt in range(4):
        try:
            settle_detail(page)
            if _otp_field_visible(page):
                log("  OTP field already on page - Action already selected")
                return True
            if _wait_button(page, r"generate otp", 1500):
                log("  GENERATE OTP already visible - Action already selected")
                return True
            cur = _action_current_text(page)
            log(f"  Action box now: {cur[:80]}")
            if want == "return" and re.search(r"return\s+to", cur, re.I):
                log("  Action already Return")
                if _wait_button(page, r"generate otp", 5000) or _otp_field_visible(page):
                    return True
            if want == "forward" and re.search(r"forward\s+to", cur, re.I):
                log("  Action already Forward")
                if _wait_button(page, r"generate otp", 5000) or _otp_field_visible(page):
                    return True
            _scroll_action_into_view(page)
            if _native_action_select(page, want):
                if _wait_button(page, r"generate otp", 8000) or _otp_field_visible(page):
                    log("  GENERATE OTP / OTP field appeared")
                    return True
            if not _click_select_one(page):
                last = "Select One not clickable"
                continue
            page.wait_for_timeout(400)
            if not _pick_from_open_list(page, want):
                last = "list option not selected"
                save_debug(page, "action_overlay_fail")
                continue
            page.wait_for_timeout(500)
            if _wait_button(page, r"generate otp", 8000) or _otp_field_visible(page):
                log("  GENERATE OTP / OTP field appeared")
                return True
            last = "GENERATE OTP not visible"
            save_debug(page, "action_no_otp_btn")
        except Exception as exc:
            last = str(exc)
            log(f"  Action attempt {attempt + 1} error: {exc}")
            s = str(exc).lower()
            if "destroyed" in s or "navigation" in s:
                page.wait_for_timeout(1000)
                continue
    log(f"  Action failed: {last}")
    save_debug(page, "action_select_fail")
    return False


def _wait_button(page, pattern: str, ms: int = 8000) -> bool:
    deadline = datetime.now().timestamp() + ms / 1000
    rx = re.compile(pattern, re.I)
    while datetime.now().timestamp() < deadline:
        try:
            btn = page.get_by_role("button", name=rx)
            for i in range(btn.count()):
                el = btn.nth(i)
                if el.is_visible():
                    box = el.bounding_box()
                    if box and box["width"] > 20:
                        return True
        except Exception:
            pass
        try:
            found = page.evaluate(
                """(pat) => {
                    const re = new RegExp(pat, 'i');
                    return [...document.querySelectorAll('button, a.btn, .btn')].some(b => {
                        const r = b.getBoundingClientRect();
                        return r.width > 20 && r.height > 10 && re.test((b.innerText || '').trim());
                    });
                }""",
                pattern,
            )
            if found:
                return True
        except Exception:
            pass
        page.wait_for_timeout(200)
    return False


def generate_otp_and_submit(page, btype: str) -> bool:
    if _otp_field_visible(page):
        log("  OTP field already visible - skip GENERATE OTP click")
    elif not _wait_button(page, r"generate otp", 10000):
        log("  GENERATE OTP button not visible yet")
        save_debug(page, "no_generate_otp")
        return False
    else:
        try:
            btn = page.get_by_role("button", name=re.compile("generate otp", re.I))
            if btn.count():
                btn.first.click(timeout=8000)
            else:
                page.evaluate(
                    """() => {
                        const b = [...document.querySelectorAll('button,a,.btn')].find(e =>
                            /generate otp/i.test(e.innerText || '')
                        );
                        if (b) b.click();
                    }"""
                )
            log("  clicked GENERATE OTP")
        except Exception as exc:
            log(f"  GENERATE OTP: {exc}")
            return False
        page.wait_for_timeout(600)
        click_modal_ok(page, wait_ms=12000)
        page.wait_for_timeout(400)
    try:
        page.keyboard.press("Escape")
    except Exception:
        pass
    page.wait_for_timeout(200)
    otp = None
    for cand in (
        page.get_by_placeholder(re.compile("otp", re.I)),
        page.locator("xpath=//*[contains(normalize-space(.),'Please Enter OTP')]/following::input[1]"),
        page.locator("input[name*='otp' i], input[id*='otp' i], input[placeholder*='OTP' i]"),
    ):
        try:
            for i in range(cand.count()):
                el = cand.nth(i)
                if el.is_visible():
                    box = el.bounding_box()
                    if box and box["width"] > 40:
                        otp = el
                        break
        except Exception:
            continue
        if otp is not None:
            break
    if otp is not None:
        _type_into(page, otp, "000000", tab=False)
        log("  OTP typed 000000")
    else:
        log("  OTP box not found (visible)")
        save_debug(page, "otp_missing")
        return False
    try:
        page.keyboard.press("Escape")
    except Exception:
        pass
    page.wait_for_timeout(400)
    # Click the footer colored button only — never the Action dropdown.
    clicked = False
    try:
        if btype == "return":
            btn = page.locator("button.btn-danger, button.btn.btn-danger").filter(
                has_text=re.compile(r"return\s+to", re.I)
            )
        else:
            btn = page.get_by_role("button", name=re.compile(r"forward\s+to", re.I))
        for i in range(btn.count()):
            el = btn.nth(i)
            try:
                if not el.is_visible():
                    continue
                box = el.bounding_box()
                if not box or box["height"] < 20:
                    continue
                el.scroll_into_view_if_needed(timeout=3000)
                el.click(timeout=5000)
                log(f"  clicked footer button: {el.inner_text()[:60]}")
                clicked = True
                break
            except Exception as exc:
                log(f"  footer btn {i}: {exc}")
    except Exception as exc:
        log(f"  footer locate: {exc}")
    if not clicked:
        clicked_js = page.evaluate(
            """(want) => {
                const btns = [...document.querySelectorAll('button')].filter(b => {
                    const t = (b.innerText || '').replace(/\\s+/g, ' ').trim();
                    const r = b.getBoundingClientRect();
                    if (r.width < 40 || r.height < 20) return false;
                    if (b.closest('ng-select, .ng-select, .ng-dropdown-panel, .cdk-overlay-pane')) return false;
                    if (/save as draft|generate otp|attach digital|upload|select one/i.test(t)) return false;
                    const cls = (b.className || '').toString();
                    if (want === 'return') return /return\\s+to/i.test(t) && /danger/.test(cls);
                    return /forward\\s+to/i.test(t);
                });
                if (!btns.length) return {ok: false};
                const t = (btns[0].innerText || '').replace(/\\s+/g, ' ').trim();
                btns[0].click();
                return {ok: true, text: t};
            }""",
            "return" if btype == "return" else "forward",
        )
        log(f"  clicked final action js: {clicked_js}")
        clicked = bool(clicked_js and clicked_js.get("ok"))
    if not clicked:
        save_debug(page, "final_action_fail")
        return False
    page.wait_for_timeout(600)
    click_modal_ok(page, wait_ms=8000)
    try:
        page.wait_for_url(re.compile(r"/success/"), timeout=20000)
        log(f"  success page: {page.url}")
    except Exception:
        log(f"  waiting success page - now {page.url}")
        click_modal_ok(page, wait_ms=3000)
        try:
            page.wait_for_url(re.compile(r"/success/"), timeout=10000)
            log(f"  success page: {page.url}")
        except Exception:
            save_debug(page, "no_success_page")
            return False
    page.wait_for_timeout(400)
    return True


def load_batch(folder: Path) -> list[dict]:
    """Batch folder: Upload Template.xlsx + PDFs named '{Letter} TPI Report...'."""
    folder = Path(folder)
    if not folder.exists():
        log(f"Batch folder not found: {folder}")
        return []
    xlsx = None
    for p in folder.glob("*.xlsx"):
        if p.name.startswith("~$"):
            continue
        xlsx = p
        if "upload" in p.name.lower() or "template" in p.name.lower():
            break
    if not xlsx:
        log("No Excel (Upload Template) in batch folder")
        return []
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(max_row=500, max_col=16, values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [re.sub(r"\s+", " ", str(c or "")).strip().lower() for c in rows[0]]

    def col(*keys):
        for i, h in enumerate(headers):
            if any(k in h for k in keys):
                return i
        return None

    isid = col("scheme")
    imb = col("mb")
    iamt = col("amount")
    ilet = col("letter")
    irem = col("remark")
    istat = col("status", "bill type", "action")
    log(f"Batch excel: {xlsx.name} headers={headers}")
    bills = []
    for raw in rows[1:]:
        if not raw or all(v is None or str(v).strip() == "" for v in raw):
            continue
        def cell(i):
            if i is None or i >= len(raw):
                return ""
            return str(raw[i] or "").strip()
        sid = cell(isid)
        mb = cell(imb)
        letter = cell(ilet)
        status = cell(istat)
        btype = _norm_bill_type(status)
        if status.lower().startswith("verif") or status.lower() == "forward":
            btype = "forward"
        amount = "0.00" if btype == "return" else cell(iamt)
        files = []
        if letter:
            pat = re.compile(rf"^{re.escape(letter)}(?:[\s_\-]|$)", re.I)
            for p in sorted(folder.iterdir()):
                if not p.is_file() or p.name.startswith("~$"):
                    continue
                if p.suffix.lower() in {".xlsx", ".xls"}:
                    continue
                if pat.match(p.name):
                    files.append(str(p))
        bills.append({
            "scheme_id": sid,
            "mb_no": mb.replace("-", "/"),
            "district": "",
            "path": str(folder),
            "files": files,
            "n_files": len(files),
            "fill": {
                "amount": amount,
                "letter": letter,
                "remark": cell(irem),
                "bill_type": btype,
                "status": status,
            },
        })
    return bills


TEMPLATE_HEADERS = [
    "Scheme ID",
    "MB No.",
    "Amount Approved by TPI",
    "TPI Letter Number",
    "Remark",
    "Bill Type",
]


def write_upload_template(path: str | Path) -> Path:
    """Create Upload Template.xlsx the operator fills, then Scan Excel + files."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Upload"
    head_fill = PatternFill("solid", fgColor="14213D")
    head_font = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
    hint_fill = PatternFill("solid", fgColor="FFF4E6")
    for i, h in enumerate(TEMPLATE_HEADERS, 1):
        cell = ws.cell(1, i, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.cell(2, 1, "20130179")
    ws.cell(2, 2, "RA/CIVIL/001")
    ws.cell(2, 3, "123456.00")
    ws.cell(2, 4, "TPI/CKT/001")
    ws.cell(2, 5, "Verified and recommended for payment")
    ws.cell(2, 6, "Forward")
    for c in range(1, 7):
        ws.cell(2, c).fill = hint_fill
    dv = DataValidation(type="list", formula1='"Forward,Return"', allow_blank=False)
    dv.error = "Use Forward or Return"
    dv.errorTitle = "Bill Type"
    dv.prompt = "Forward = Verified (DSC). Return = amount 0.00, no DSC."
    dv.promptTitle = "Bill Type"
    ws.add_data_validation(dv)
    dv.add("F2:F200")
    note = wb.create_sheet("How to use")
    lines = [
        "TPI Upload Template",
        "1. Delete the sample row (row 2) before a real batch.",
        "2. One row = one bill. Scheme ID and MB No. must match the portal list.",
        "3. Bill Type: Forward or Return. Return sets Amount to 0.00 and skips DSC.",
        "4. TPI Letter Number is also the PDF prefix. File name must start with that letter, e.g. TPI/CKT/001 TPI Report.pdf",
        "5. Put this Excel and all PDFs in the same batch folder. Each PDF max 7 MB.",
        "6. In the app: Browse that folder → Scan Excel + files → PROCEED.",
        "7. Parallel: the app logs in once, then opens up to 5 windows for 5 bills, then the next 5.",
    ]
    for i, t in enumerate(lines, 1):
        note.cell(i, 1, t)
        note.cell(i, 1).font = Font(bold=(i == 1), name="Calibri", size=12 if i == 1 else 11)
    note.column_dimensions["A"].width = 110
    widths = [16, 16, 26, 22, 48, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:F200"
    wb.save(path)
    return path


def upload_one(page, b: dict, steps: set[str]) -> None:
    fill = b.get("fill") or {}
    btype = _norm_bill_type(fill.get("bill_type") or fill.get("status"))
    row = {"scheme_id": b["scheme_id"], "mb_no": b["mb_no"]}
    go_to_tpi_list(page)
    detail = open_row(page, row)
    grand = scrape_grand_total(detail)
    if grand:
        log(f"  Grand Total: {grand}")
    if btype == "return":
        fill["amount"] = "0.00"
    log(f"  bill type: {btype}  steps={sorted(steps)}")
    if "files" in steps:
        uploaded = upload_files_on_detail(detail, b["files"])
        log(f"  files uploaded: {len(uploaded)}")
    if "fill" in steps:
        fill_tpi_fields(detail, fill)
        saved = click_save_as_draft(detail)
        if not saved:
            log("  retry fill letter + save")
            fill_tpi_fields(detail, fill)
            saved = click_save_as_draft(detail)
        log("  waiting for Action box after Save Draft...")
        settle_detail(detail)
    if "dsc" in steps and btype == "forward":
        click_attach_signature(detail)
    elif "dsc" in steps:
        log("  Return bill — skip DSC")
    ok_act = False
    if "action" in steps:
        try:
            settle_detail(detail)
            if select_action(detail, btype):
                ok_act = generate_otp_and_submit(detail, btype)
            else:
                log("  skip OTP — action not selected")
        except Exception as exc:
            log(f"  Action/OTP step: {exc}")
            save_debug(detail, "action_step_fail")
        try:
            batch = Path(b.get("path") or os.getenv("TPI_BATCH") or ".")
            tag = "Return" if btype == "return" else "Forward"
            st = "OK" if ok_act else "FAIL"
            fn = f"{b.get('scheme_id')}_{win_name(b.get('mb_no') or '', 20)}_{tag}_{st}_{datetime.now().strftime('%H%M%S')}.png"
            dest = batch / fn
            detail.screenshot(path=str(dest), full_page=True)
            log(f"  screenshot saved: {dest.name}")
        except Exception as exc:
            log(f"  screenshot: {exc}")
    try:
        from activity import log_event
        log_event(
            kind="upload",
            action="return" if btype == "return" else "forward",
            cluster=os.getenv("TPI_CLUSTER") or "",
            district=b.get("district") or "",
            scheme_id=b.get("scheme_id") or "",
            mb_no=b.get("mb_no") or "",
            amount=(fill or {}).get("amount") or "",
            grand_total=grand,
            letter=(fill or {}).get("letter") or "",
        )
    except Exception as exc:
        log(f"  activity: {exc}")
    if detail != page:
        try:
            detail.close()
        except Exception:
            pass
    try:
        back_to_tpi_list(page)
    except Exception:
        go_to_tpi_list(page)


def _worker_browser(bill: dict, steps: set[str], state_path: str, headed: bool, idx: int, total: int) -> str:
    sid = bill.get("scheme_id") or ""
    mb = bill.get("mb_no") or ""
    log(f"[{idx}/{total}] parallel window  {sid} {mb}")
    from playwright.sync_api import sync_playwright
    try:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(headless=not headed)
            context = browser.new_context(storage_state=state_path, accept_downloads=True)
            page = context.new_page()
            try:
                go_to_tpi_list(page)
                upload_one(page, bill, steps)
                return f"OK {sid} {mb}"
            finally:
                context.close()
                browser.close()
    except Exception as exc:
        log(f"  ERROR {sid} {mb}: {exc}")
        return f"FAIL {sid} {mb}: {exc}"


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="TPI MB uploader")
    p.add_argument("--headed", action="store_true")
    p.add_argument("--scan-only", action="store_true")
    p.add_argument("--batch", default="", help="Folder with Upload Template.xlsx + TPI Report PDFs")
    p.add_argument("--day", default="", help="DD.MM.YYYY (legacy)")
    p.add_argument("--session", default="")
    p.add_argument("--skip-login", action="store_true")
    p.add_argument("--parallel", type=int, default=0, help="Bills at once after login (default 5, max 10)")
    return p.parse_args()


def main() -> None:
    args = parse_args()
    batch = (args.batch or os.getenv("TPI_BATCH") or "").strip()
    if batch:
        log(f"Batch folder: {batch}")
        bills = load_batch(Path(batch))
    else:
        root = DOWNLOAD_DIR
        day = args.day.strip() or today_stamp()
        session = args.session.strip()
        if not session:
            sess = list_sessions(root, day)
            session = sess[-1] if sess else "Session 1"
        log(f"Upload root: {root}")
        log(f"Day: {day}  Session: {session}")
        bills = scan_session(root, day, session)
    pending = [
        b for b in bills
        if b["n_files"] or (b.get("fill") or {}).get("letter") or (b.get("fill") or {}).get("amount")
    ]
    only = (os.getenv("TPI_ONLY") or "").strip()
    if only:
        keys = {k.strip() for k in only.split(";") if k.strip()}
        pending = [b for b in pending if f"{b.get('scheme_id')}|{b.get('mb_no')}" in keys]
        log(f"Filtered to selected rows: {len(pending)}")
    log(f"Rows: {len(bills)} | ready: {len(pending)}")
    for b in bills:
        fill = b.get("fill") or {}
        log(f"  {b.get('scheme_id')} | {b.get('mb_no')} | letter={fill.get('letter')} | {fill.get('bill_type')} | files={b['n_files']}")
        for f in b["files"]:
            log(f"      {Path(f).name}")

    if args.scan_only:
        log("Scan only. No portal upload.")
        return
    if not pending:
        log("Nothing to upload. Create For Upload\\ with files and/or an Excel (Amount, Letter, Remark).")
        return
    steps = {s.strip() for s in (os.getenv("TPI_STEPS") or "files,fill,dsc,action").split(",") if s.strip()}
    log(f"Steps: {sorted(steps)}")
    if not USER or not PASSWORD:
        sys.exit("ERROR: PORTAL_USER / PORTAL_PASS missing")

    try:
        npar = int(args.parallel or os.getenv("TPI_PARALLEL") or 5)
    except Exception:
        npar = 5
    npar = max(1, min(10, npar))
    log(f"Parallel windows after login: {npar}")

    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=not args.headed)
        context = browser.new_context(accept_downloads=True)
        page = context.new_page()
        try:
            if not args.skip_login:
                login(page)
            go_to_tpi_list(page)
            if npar <= 1 or len(pending) <= 1:
                for i, b in enumerate(pending, 1):
                    log(f"[{i}/{len(pending)}] upload {b['scheme_id']} {b['mb_no']}")
                    try:
                        upload_one(page, b, steps)
                    except Exception as exc:
                        log(f"  ERROR: {exc}")
                        save_debug(page, f"upload_err_{b['scheme_id']}")
                        try:
                            go_to_tpi_list(page)
                        except Exception:
                            pass
            else:
                fd, state_path = tempfile.mkstemp(suffix=".json")
                os.close(fd)
                context.storage_state(path=state_path)
                log("Login saved. Opening parallel windows (CAPTCHA not needed again).")
                try:
                    for start in range(0, len(pending), npar):
                        wave = pending[start : start + npar]
                        wnum = start // npar + 1
                        wtot = (len(pending) + npar - 1) // npar
                        log(f"======== Wave {wnum}/{wtot}  {len(wave)} window(s) ========")
                        with ThreadPoolExecutor(max_workers=len(wave)) as pool:
                            futs = [
                                pool.submit(
                                    _worker_browser,
                                    b,
                                    steps,
                                    state_path,
                                    args.headed,
                                    start + i + 1,
                                    len(pending),
                                )
                                for i, b in enumerate(wave)
                            ]
                            for fut in as_completed(futs):
                                log(f"  wave result: {fut.result()}")
                finally:
                    try:
                        os.remove(state_path)
                    except Exception:
                        pass
        finally:
            context.close()
            browser.close()
    log("Upload finished.")


if __name__ == "__main__":
    main()
